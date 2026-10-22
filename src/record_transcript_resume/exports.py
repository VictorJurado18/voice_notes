"""Export saved Voice Notes documents without requiring an external editor."""
from __future__ import annotations

from html import escape
from pathlib import Path
import re
import tempfile
import time
from uuid import uuid4


def export_docx(source: Path, destination: Path) -> Path:
    """Export Markdown-like text as an editable Word document."""
    from docx import Document

    document = Document()
    for line in source.read_text(encoding="utf-8").splitlines():
        if line.startswith("# "):
            document.add_heading(line[2:].strip(), level=0)
        elif line.startswith("## "):
            document.add_heading(line[3:].strip(), level=1)
        elif line.strip():
            document.add_paragraph(line)
    destination.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(destination))
    return destination


def export_pdf(source: Path, destination: Path, *, content: str | None = None) -> Path:
    """Export text as a readable PDF with Portuguese characters preserved."""
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.units import cm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

    styles = getSampleStyleSheet()
    story = []
    text = source.read_text(encoding="utf-8") if content is None else content
    for line in text.splitlines():
        if line.startswith("# "):
            story.append(Paragraph(escape(line[2:].strip()), styles["Title"]))
        elif line.startswith("## "):
            story.append(Spacer(1, 0.25 * cm))
            story.append(Paragraph(escape(line[3:].strip()), styles["Heading2"]))
        elif line.strip():
            story.append(Paragraph(escape(line).replace("\n", "<br/>"), styles["BodyText"]))
            story.append(Spacer(1, 0.1 * cm))
    destination.parent.mkdir(parents=True, exist_ok=True)
    SimpleDocTemplate(str(destination), pagesize=A4, rightMargin=1.8 * cm, leftMargin=1.8 * cm).build(story)
    return destination


def create_pdf_preview(source: Path, *, content: str | None = None) -> Path:
    """Keep a temporary PDF available for the external viewer's Save As action.

    Each preview has a private directory so it can stay open independently.
    Only our old UUID directories are pruned; an open/locked file is retained.
    """
    root = Path(tempfile.gettempdir()) / "VoiceNotes" / "pdf-previews"
    root.mkdir(parents=True, exist_ok=True)
    if root.is_symlink() or root.is_junction():
        raise OSError("A pasta temporária de PDFs não pode ser um link ou junção.")
    resolved_root = root.resolve(strict=True)
    cutoff = time.time() - 7 * 24 * 60 * 60
    for folder in root.iterdir():
        try:
            if (not re.fullmatch(r"[0-9a-f]{32}", folder.name)
                    or folder.is_symlink() or folder.is_junction() or not folder.is_dir()
                    or folder.resolve().parent != resolved_root
                    or folder.stat().st_mtime >= cutoff):
                continue
            for preview in folder.iterdir():
                if preview.is_file() and not preview.is_symlink() and preview.suffix == ".pdf":
                    preview.unlink()
            folder.rmdir()
        except OSError:
            continue
    folder = root / uuid4().hex
    folder.mkdir()
    destination = folder / f"{source.stem[:80] or 'documento'}.pdf"
    try:
        return export_pdf(source, destination, content=content)
    except Exception:
        try:
            destination.unlink(missing_ok=True)
            folder.rmdir()
        except OSError:
            pass
        raise
